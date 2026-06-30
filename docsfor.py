from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Set margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Shop Header
header = doc.add_paragraph()
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = header.add_run('ANAND SALE')
run.bold = True
run.font.size = Pt(20)

address = doc.add_paragraph()
address.alignment = WD_ALIGN_PARAGRAPH.CENTER
address.add_run('Near Trishna Resturant, Harmu Road, Argora, Ranchi-834002, Jharkhand\n').font.size = Pt(11)
address.add_run('Phone: +91-9199681627 | Email: anandsales006@gmail.com\n').font.size = Pt(10)
gst_run = address.add_run('GSTIN: 20ABCDE1234F1Z5')
gst_run.bold = True
gst_run.font.size = Pt(10)

doc.add_paragraph('_' * 85).alignment = WD_ALIGN_PARAGRAPH.CENTER

# Invoice Meta
meta = doc.add_paragraph()
meta.add_run('TAX INVOICE').bold = True
meta.add_run('\nInvoice No: AS-2024-001\t\t\tDate: 18/12/2024')

# Customer Details
cust = doc.add_paragraph()
cust.add_run('Bill To:\n').bold = True
cust.add_run('Customer Name: Ankesh Kumar\n')
cust.add_run('Phone Number: 8092720126\n')
cust.add_run('Address: Kishoreganj Chowk, Ranchi')

doc.add_paragraph('_' * 85).alignment = WD_ALIGN_PARAGRAPH.CENTER

# Table Data
table = doc.add_table(rows=2, cols=9)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ['S.No', 'Item Description & Details', 'HSN', 'Qty', 'Rate', 'Taxable Value', 'CGST 9%', 'SGST 9%', 'Total']
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            run.bold = True

row_cells = table.rows[1].cells
row_cells[0].text = '1'
row_cells[1].text = ('Samsung Galaxy S23 Ultra\n'
                     '• IMEI 1: 353835880765880 / 01\n'
                     '• IMEI 2: 353901650765888 / 01\n'
                     '• EID: 8904305120220000622200996961685\n'
                     '• S/N: RZCW217N2SD')
row_cells[2].text = '8517'
row_cells[3].text = '1'
row_cells[4].text = '90,677.97'
row_cells[5].text = '90,677.97'
row_cells[6].text = '8,161.02'
row_cells[7].text = '8,161.01'
row_cells[8].text = '1,07,000.00'

# Total Amount
doc.add_paragraph()
total_p = doc.add_paragraph()
total_p.add_run('Total Amount (in words): ').bold = True
total_p.add_run('One Lakh Seven Thousand Rupees Only.')

total_val = doc.add_paragraph()
total_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = total_val.add_run('Total Invoice Value: ₹ 1,07,000.00')
run.bold = True
run.font.size = Pt(12)

doc.add_paragraph('_' * 85).alignment = WD_ALIGN_PARAGRAPH.CENTER

# Terms and Signatory
terms = doc.add_paragraph()
terms.add_run('Terms & Conditions:\n').bold = True
terms.add_run('1. Goods once sold will not be taken back.\n')
terms.add_run('2. Warranty is subject to manufacturer terms & conditions.\n')
terms.add_run('3. Please retain this invoice for warranty claims.\n\n')

sign = doc.add_paragraph()
sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
sign.add_run('For Anand Sale:\n\n\n')
sign.add_run('(Authorised Signatory)').bold = True

# Save Document
doc.save('Anand_Sale_Invoice.docx')
print("Invoice generated successfully: Anand_Sale_Invoice.docx")